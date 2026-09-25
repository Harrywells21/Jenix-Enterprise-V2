"""Tests for routes/cve.py: CVE scan trigger (background task + DB
persistence), results (cache-then-DB fallback), fleet-wide summary
(cache+DB merge), and the Excel export's dual auth path.

_scan_cache is a process-lifetime module-level dict in routes/cve.py --
it is NOT reset by the per-test db_session/Base.metadata teardown, so an
autouse fixture below clears it before AND after every test in this file
to prevent state leaking across tests (see the continuity doc's standing
rule on this).

_do_scan() (the background task fired by POST /scan/{machine_id}) opens
its own DB session via SessionLocal() imported directly from db.py --
it does NOT go through the get_db dependency the `client` fixture
overrides. Tests that trigger a real scan therefore monkeypatch
routes.cve.SessionLocal to a sessionmaker bound to conftest's shared
test engine, so the background task's writes land in the same
in-memory DB the test itself queries afterward -- never the real
server/jenix.db.
"""
import io
from datetime import datetime

import pytest
from openpyxl import load_workbook

from conftest import auth_headers, engine
from sqlalchemy.orm import sessionmaker

from db import Machine, CveScan, CveFinding, AuditLog
from auth import create_token
import routes.cve as cve_module


@pytest.fixture(autouse=True)
def _clear_scan_cache():
    cve_module._scan_cache.clear()
    yield
    cve_module._scan_cache.clear()


def _make_machine(db_session, hostname="test-machine", token="tok-cve-1"):
    m = Machine(hostname=hostname, ip="10.0.0.60", os_name="Linux",
                kernel="6.8", token=token, status="online")
    db_session.add(m)
    db_session.commit()
    db_session.refresh(m)
    return m


def _canned_scan_result():
    return {
        "packages_scanned": 2,
        "vulnerable_packages": 1,
        "total_vulns": 2,
        "critical": 0,
        "high": 1,
        "risk_level": "HIGH",
        "results": [
            {
                "package": "openssl",
                "version": "3.0.2-0ubuntu1",
                "vulns": [
                    {"id": "CVE-TEST-0001", "summary": "Test vuln one",
                     "severity": "HIGH",
                     "url": "https://osv.dev/vulnerability/CVE-TEST-0001"},
                    {"id": "CVE-TEST-0002", "summary": "Test vuln two",
                     "severity": "LOW",
                     "url": "https://osv.dev/vulnerability/CVE-TEST-0002"},
                ],
                "count": 2,
                "highest": "HIGH",
            }
        ],
    }


# ── POST /cve/scan/{machine_id} ──────────────────────────────────────────

def test_trigger_scan_404_unknown_machine(client, db_session, operator_user):
    h = auth_headers(operator_user)
    r = client.post("/api/cve/scan/999999", headers=h)
    assert r.status_code == 404


def test_trigger_scan_requires_operator_not_viewer(client, db_session, admin_user, viewer_user):
    m = _make_machine(db_session, hostname="scan-viewer-blocked", token="tok-scan-viewer")
    h = auth_headers(viewer_user)
    r = client.post(f"/api/cve/scan/{m.id}", headers=h)
    assert r.status_code == 403


def test_trigger_scan_allows_operator(client, db_session, admin_user, operator_user, monkeypatch):
    m = _make_machine(db_session, hostname="scan-operator-allowed", token="tok-scan-operator")
    TestSessionLocal = sessionmaker(bind=engine, autocommit=False, autoflush=False)
    monkeypatch.setattr(cve_module, "SessionLocal", TestSessionLocal)
    monkeypatch.setattr(cve_module, "run_cve_scan", lambda max_packages=30: _canned_scan_result())
    h = auth_headers(operator_user)
    r = client.post(f"/api/cve/scan/{m.id}", headers=h)
    assert r.status_code == 200


def test_trigger_scan_creates_audit_log(client, db_session, admin_user, operator_user, monkeypatch):
    m = _make_machine(db_session, hostname="scan-audit", token="tok-scan-audit")
    TestSessionLocal = sessionmaker(bind=engine, autocommit=False, autoflush=False)
    monkeypatch.setattr(cve_module, "SessionLocal", TestSessionLocal)
    monkeypatch.setattr(cve_module, "run_cve_scan", lambda max_packages=30: _canned_scan_result())
    h = auth_headers(operator_user)
    client.post(f"/api/cve/scan/{m.id}", headers=h)
    log = db_session.query(AuditLog).filter(
        AuditLog.machine_id == m.id, AuditLog.action == "cve_scan").first()
    assert log is not None
    assert log.status == "ok"


def test_trigger_scan_populates_cache_and_db(client, db_session, admin_user, operator_user, monkeypatch):
    m = _make_machine(db_session, hostname="scan-target", token="tok-scan-1")
    TestSessionLocal = sessionmaker(bind=engine, autocommit=False, autoflush=False)
    monkeypatch.setattr(cve_module, "SessionLocal", TestSessionLocal)
    monkeypatch.setattr(cve_module, "run_cve_scan", lambda max_packages=30: _canned_scan_result())

    h = auth_headers(operator_user)
    r = client.post(f"/api/cve/scan/{m.id}", headers=h)
    assert r.status_code == 200
    assert "started" in r.json()["message"].lower()

    assert m.id in cve_module._scan_cache
    cached = cve_module._scan_cache[m.id]
    assert cached["risk_level"] == "HIGH"
    assert "scanned_at" in cached

    scan_row = db_session.query(CveScan).filter(
        CveScan.machine_id == m.id).order_by(CveScan.scanned_at.desc()).first()
    assert scan_row is not None
    assert scan_row.risk_level == "HIGH"
    findings = db_session.query(CveFinding).filter(
        CveFinding.scan_id == scan_row.id).all()
    assert len(findings) == 2


# ── GET /cve/results/{machine_id} ────────────────────────────────────────

def test_results_no_scan_yet(client, db_session, admin_user):
    m = _make_machine(db_session, hostname="results-empty", token="tok-results-empty")
    h = auth_headers(admin_user)
    r = client.get(f"/api/cve/results/{m.id}", headers=h)
    assert r.status_code == 200
    assert r.json()["scanned"] is False


def test_results_reads_from_cache(client, db_session, admin_user):
    m = _make_machine(db_session, hostname="results-cache", token="tok-results-cache")
    cve_module._scan_cache[m.id] = {
        "scanned_at": "2026-01-01T00:00:00",
        "packages_scanned": 1, "vulnerable_packages": 1, "total_vulns": 1,
        "critical": 0, "high": 0, "risk_level": "LOW", "results": [],
    }
    h = auth_headers(admin_user)
    r = client.get(f"/api/cve/results/{m.id}", headers=h)
    assert r.status_code == 200
    body = r.json()
    assert body["scanned"] is True
    assert body["source"] == "cache"


def test_results_falls_back_to_db_when_cache_empty(client, db_session, admin_user):
    m = _make_machine(db_session, hostname="results-db", token="tok-results-db")
    scan = CveScan(machine_id=m.id, triggered_by_id=admin_user.id,
                    triggered_by_name=admin_user.name, scanned_at=datetime.utcnow(),
                    packages_scanned=1, vulnerable_packages=1, total_vulns=2,
                    critical=0, high=1, risk_level="HIGH")
    db_session.add(scan); db_session.commit(); db_session.refresh(scan)
    f1 = CveFinding(scan_id=scan.id, package="openssl", version="3.0.2",
                     cve_id="CVE-X-1", summary="s", severity="HIGH",
                     url="https://osv.dev/x1")
    f2 = CveFinding(scan_id=scan.id, package="openssl", version="3.0.2",
                     cve_id="CVE-X-2", summary="s2", severity="LOW",
                     url="https://osv.dev/x2")
    db_session.add_all([f1, f2]); db_session.commit()

    h = auth_headers(admin_user)
    r = client.get(f"/api/cve/results/{m.id}", headers=h)
    assert r.status_code == 200
    body = r.json()
    assert body["source"] == "db"
    assert body["risk_level"] == "HIGH"
    assert len(body["results"]) == 1
    assert body["results"][0]["package"] == "openssl"
    assert len(body["results"][0]["vulns"]) == 2


# ── GET /cve/summary ──────────────────────────────────────────────────────

def test_summary_empty_state(client, db_session, admin_user):
    h = auth_headers(admin_user)
    r = client.get("/api/cve/summary", headers=h)
    assert r.status_code == 200
    body = r.json()
    assert body["machines_scanned"] == 0
    assert body["last_scans"] == {}


def test_summary_merges_cache_and_db_viewer_allowed(client, db_session, admin_user, viewer_user):
    m1 = _make_machine(db_session, hostname="summary-cache", token="tok-summary-cache")
    m2 = _make_machine(db_session, hostname="summary-db", token="tok-summary-db")

    cve_module._scan_cache[m1.id] = {
        "scanned_at": "2026-02-01T00:00:00", "critical": 1, "high": 0,
        "vulnerable_packages": 1, "risk_level": "CRITICAL",
    }
    scan = CveScan(machine_id=m2.id, triggered_by_id=admin_user.id,
                    triggered_by_name=admin_user.name, scanned_at=datetime.utcnow(),
                    packages_scanned=1, vulnerable_packages=1, total_vulns=1,
                    critical=0, high=1, risk_level="HIGH")
    db_session.add(scan); db_session.commit()

    h = auth_headers(viewer_user)
    r = client.get("/api/cve/summary", headers=h)
    assert r.status_code == 200
    body = r.json()
    assert body["machines_scanned"] == 2
    assert body["total_critical"] == 1
    assert body["total_high"] == 1
    assert body["last_scans"][str(m1.id)]["source"] == "cache"
    assert body["last_scans"][str(m2.id)]["source"] == "db"


# ── GET /cve/export/excel ────────────────────────────────────────────────

def test_export_excel_with_bearer_header(client, db_session, admin_user):
    h = auth_headers(admin_user)
    r = client.get("/api/cve/export/excel", headers=h)
    assert r.status_code == 200
    assert r.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == [
        "Executive Summary", "Full CVE Detail",
        "Scan Activity Log", "Remediation Guidance",
    ]
    assert wb["Executive Summary"]["A1"].value == (
        "JENIX Enterprise — CVE Security Report")


def test_export_excel_with_query_token(client, db_session, admin_user):
    token = create_token({"sub": str(admin_user.id)})
    r = client.get(f"/api/cve/export/excel?token={token}")
    assert r.status_code == 200
    assert r.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def test_export_excel_401_with_no_token(client, db_session):
    r = client.get("/api/cve/export/excel")
    assert r.status_code == 401


def test_export_excel_content_reflects_real_data(client, db_session, admin_user):
    m = _make_machine(db_session, hostname="export-data", token="tok-export-data")
    scan = CveScan(machine_id=m.id, triggered_by_id=admin_user.id,
                    triggered_by_name=admin_user.name, scanned_at=datetime.utcnow(),
                    packages_scanned=1, vulnerable_packages=1, total_vulns=1,
                    critical=0, high=1, risk_level="HIGH")
    db_session.add(scan); db_session.commit(); db_session.refresh(scan)
    f = CveFinding(scan_id=scan.id, package="nginx", version="1.24.0",
                    cve_id="CVE-EXPORT-1", summary="export test finding",
                    severity="HIGH",
                    url="https://osv.dev/vulnerability/CVE-EXPORT-1")
    db_session.add(f); db_session.commit()

    h = auth_headers(admin_user)
    r = client.get("/api/cve/export/excel", headers=h)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    detail_ws = wb["Full CVE Detail"]
    values = [cell.value for cell in detail_ws[2]]
    assert "nginx" in values
    assert "CVE-EXPORT-1" in values
